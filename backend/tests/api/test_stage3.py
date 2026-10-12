"""Stage 3: wallet & store (R20, R21, R25), finance (R40), tickets (R32, R33, R39), reports (R31, R41),
broadcasts (R46), surveys (R23)."""

import io
from datetime import UTC, datetime

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core import clock
from app.modules.notifications.models import Notification, SmsOutbox
from app.modules.points import service as points
from app.modules.points.models import PointsLedger
from tests.conftest import jpeg_bytes, make_staff, make_user, staff_headers, user_headers


def _give_points(db, u, amount, key="seed"):
    """Credit `amount` AVAILABLE points (as if the month had already ended)."""
    e = points.award(db, user_id=u.id, event_type="signup", idempotency_key=f"{key}:{u.id}")
    e.status, e.amount = "available", amount
    db.commit()


def test_month_end_release(client, db):
    u = make_user(db)
    points.award(db, user_id=u.id, event_type="signup", idempotency_key="s")
    db.commit()
    assert points.balance(db, u.id).pending == 100
    clock.freeze(datetime(2026, 10, 22, 20, 29, tzinfo=UTC))  # one minute before 1 Aban 00:00 Tehran
    assert points.release_due(db) == 0
    clock.freeze(datetime(2026, 10, 22, 20, 30, tzinfo=UTC))
    assert points.release_due(db) == 1 and points.release_due(db) == 0
    db.commit()
    assert points.balance(db, u.id).available == 100


def _store(client, u):
    return client.get("/api/v1/app/store", headers=user_headers(u)).json()


def test_store_categories_match_requirement(client, db):
    u = make_user(db)
    db.commit()
    keys = {c["key"] for c in _store(client, u)["categories"]}
    assert keys == {"cash", "pet_food", "insulation", "repair_discount", "charity", "stationery", "gift"}  # R25


def _item(client, fin, kind="product", cost=300, stock=None, cat_key="insulation", **kw):
    cats = {c["key"]: c["id"] for c in client.get("/api/v1/admin/store/categories", headers=staff_headers(fin)).json()}
    body = {"category_id": cats[cat_key], "title": kw.pop("title", "بستهٔ درزگیر پنجره"), "kind": kind, "points_cost": cost,
            "stock": stock, **kw}
    r = client.post("/api/v1/admin/store/items", json=body, headers=staff_headers(fin))
    assert r.status_code == 201, r.text
    return r.json()


def test_redeem_product_reserve_and_reject_refunds(client, db):
    u = make_user(db)
    fin = make_staff(db, "finance")
    db.commit()
    _give_points(db, u, 500)
    item = _item(client, fin, stock=1, needs_address=True)
    assert client.post("/api/v1/app/store/redeem", json={"item_id": item["id"]}, headers=user_headers(u)).status_code == 422  # address
    r = client.post("/api/v1/app/store/redeem", json={"item_id": item["id"], "delivery_address": "رشت، گلسار"}, headers=user_headers(u))
    assert r.status_code == 201
    rid = r.json()["id"]
    assert _store(client, u)["available_points"] == 200  # reserved immediately
    out = client.post("/api/v1/app/store/redeem", json={"item_id": item["id"], "delivery_address": "x"}, headers=user_headers(u))
    assert out.json()["error"]["code"] == "out_of_stock"
    q = client.get("/api/v1/admin/finance/redemptions", headers=staff_headers(fin)).json()
    assert q["total"] == 1 and q["items"][0]["fraud_signals"].get("new_account_days") == 0
    client.post(f"/api/v1/admin/finance/redemptions/{rid}/review", json={"approve": False, "note": "ناموجود"}, headers=staff_headers(fin))
    assert _store(client, u)["available_points"] == 500
    stock = next(i for i in _store(client, u)["items"] if i["id"] == item["id"])["stock"]
    assert stock == 1  # stock returned too


def test_cash_withdrawal_requires_iban_min_and_finance(client, db):
    u = make_user(db)
    fin = make_staff(db, "finance")
    op = make_staff(db, "operator")
    db.commit()
    _give_points(db, u, 2000)
    r = client.post("/api/v1/app/store/redeem", json={"kind": "cash", "points": 100}, headers=user_headers(u))
    assert r.status_code == 422  # below the 500 minimum
    r = client.post("/api/v1/app/store/redeem", json={"kind": "cash", "points": 1000}, headers=user_headers(u))
    assert r.json()["error"]["code"] == "iban_required"
    client.patch("/api/v1/app/me", json={"iban": "IR820540102680020817909002"}, headers=user_headers(u))
    r = client.post("/api/v1/app/store/redeem", json={"kind": "cash", "points": 1000}, headers=user_headers(u)).json()
    assert r["rial_value"] == 1_000_000 and r["status"] == "submitted"
    # operators cannot touch money (R34)
    assert client.post(f"/api/v1/admin/finance/redemptions/{r['id']}/review", json={"approve": True}, headers=staff_headers(op)).status_code == 403
    assert client.post(f"/api/v1/admin/finance/redemptions/{r['id']}/fulfill", json={"payment_ref": "x"}, headers=staff_headers(fin)).status_code == 409
    client.post(f"/api/v1/admin/finance/redemptions/{r['id']}/review", json={"approve": True}, headers=staff_headers(fin))
    assert client.post(f"/api/v1/admin/finance/redemptions/{r['id']}/fulfill", json={}, headers=staff_headers(fin)).status_code == 422
    d = client.post(f"/api/v1/admin/finance/redemptions/{r['id']}/fulfill", json={"payment_ref": "TRX-99812"}, headers=staff_headers(fin)).json()
    assert d["status"] == "fulfilled" and d["iban_masked"] == "IR82…9002"
    assert client.get("/api/v1/admin/finance/summary", headers=staff_headers(fin)).json()["cash_paid_rial"] == 1_000_000


def test_pending_points_cannot_be_spent_and_cancel(client, db):
    u = make_user(db)
    fin = make_staff(db, "finance")
    points.award(db, user_id=u.id, event_type="signup", idempotency_key="s")
    db.commit()
    item = _item(client, fin, cost=50)
    r = client.post("/api/v1/app/store/redeem", json={"item_id": item["id"]}, headers=user_headers(u))
    assert r.json()["error"]["code"] == "insufficient_points"  # pending until month end (R20)
    _give_points(db, u, 100, "x")
    rid = client.post("/api/v1/app/store/redeem", json={"item_id": item["id"]}, headers=user_headers(u)).json()["id"]
    assert client.post(f"/api/v1/app/store/redemptions/{rid}/cancel", headers=user_headers(u)).json()["status"] == "cancelled"
    assert _store(client, u)["available_points"] == 100


def test_donation_to_charity(client, db):
    u = make_user(db)
    fin = make_staff(db, "finance")
    db.commit()
    _give_points(db, u, 800)
    assert client.post("/api/v1/admin/store/items", json={"category_id": 1, "title": "x", "kind": "donation", "points_cost": 10},
                       headers=staff_headers(fin)).status_code in (404, 422)
    item = _item(client, fin, kind="donation", cost=100, cat_key="charity", title="کمک به خیریهٔ کهریزک", charity_name="کهریزک")
    r = client.post("/api/v1/app/store/redeem", json={"item_id": item["id"], "points": 600}, headers=user_headers(u)).json()
    assert r["points"] == 600 and r["kind"] == "donation"


def test_ticket_lifecycle(client, db):
    u = make_user(db)
    op = make_staff(db, "operator")
    tech = make_staff(db, "technician")
    fin = make_staff(db, "finance")
    db.commit()
    r = client.post("/api/v1/app/tickets", data={"subject_key": "reading", "title": "قرائت من تأیید نشد", "body": "لطفاً بررسی کنید",
                                                 "kinds": ["meter_photo"]},
                    files=[("files", ("m.jpg", jpeg_bytes(), "image/jpeg"))], headers=user_headers(u))
    assert r.status_code == 201, r.text
    t = r.json()
    assert t["tracking_code"].startswith("GZ-") and t["messages"][0]["attachments"][0]["kind"] == "meter_photo"
    tid = t["id"]
    # support unit sees it, finance does not (unit scoping)
    assert client.get(f"/api/v1/admin/tickets/{tid}", headers=staff_headers(op)).status_code == 200
    assert client.get(f"/api/v1/admin/tickets/{tid}", headers=staff_headers(fin)).status_code == 404
    assert any(n.kind == "ticket" for n in db.scalars(select(Notification).where(Notification.recipient_id == op.id)))
    # internal note never reaches the customer
    client.post(f"/api/v1/admin/tickets/{tid}/messages", data={"body": "به نظر عکس تار است", "internal": "true"}, headers=staff_headers(op))
    assert all(not m["is_internal"] for m in client.get(f"/api/v1/app/tickets/{tid}", headers=user_headers(u)).json()["messages"])
    # refer to the technical unit with priority + follow-up
    d = client.patch(f"/api/v1/admin/tickets/{tid}", json={"unit": "technical", "priority": "high",
                                                             "follow_up_at": "2026-10-13T06:00:00Z"}, headers=staff_headers(op)).json()
    assert d["assigned_unit"] == "technical" and d["priority"] == "high" and len(d["events"]) == 3
    assert client.get(f"/api/v1/admin/tickets/{tid}", headers=staff_headers(op)).status_code == 404  # left support's queue
    d = client.post(f"/api/v1/admin/tickets/{tid}/messages", data={"body": "لطفاً عکس واضح‌تری بفرستید", "set_status": "awaiting_customer"},
                    headers=staff_headers(tech)).json()
    assert d["status"] == "awaiting_customer"
    ut = client.get(f"/api/v1/app/tickets/{tid}", headers=user_headers(u)).json()
    assert ut["messages"][-1]["author_name"] == "پشتیبانی گازیار"  # staff identity hidden
    assert client.post(f"/api/v1/app/tickets/{tid}/rate", json={"rating": 5}, headers=user_headers(u)).status_code == 409
    client.post(f"/api/v1/app/tickets/{tid}/messages", data={"body": "ارسال شد"}, files=[("files", ("n.jpg", jpeg_bytes((9, 9, 9)), "image/jpeg"))],
                headers=user_headers(u))
    assert client.get(f"/api/v1/admin/tickets/{tid}", headers=staff_headers(tech)).json()["status"] == "open"
    client.patch(f"/api/v1/admin/tickets/{tid}", json={"status": "resolved"}, headers=staff_headers(tech))
    rated = client.post(f"/api/v1/app/tickets/{tid}/rate", json={"rating": 4, "comment": "ممنون"}, headers=user_headers(u)).json()
    assert rated["rating"] == 4 and rated["status"] == "closed" and not rated["can_reply"]
    assert client.post(f"/api/v1/app/tickets/{tid}/rate", json={"rating": 5}, headers=user_headers(u)).status_code == 409
    db.rollback()  # drop the stale read snapshot from the notification query above
    other = make_user(db, sn="7777777777", mobile="09127777777")
    db.commit()
    assert client.get(f"/api/v1/app/tickets/{tid}", headers=user_headers(other)).status_code == 404


def test_reports_run_export_and_snapshot(client, db):
    from app.modules.reports.service import REPORTS

    admin = make_staff(db)
    op = make_staff(db, "operator")
    u = make_user(db)
    u.first_name = "=HYPERLINK(\"http://evil\")"  # formula-injection probe
    db.commit()
    from app.modules.referrals.models import Referral

    db.add(Referral(referrer_user_id=u.id, referee_user_id=u.id, code_used="X", status="rewarded"))
    db.commit()
    for key in REPORTS:
        r = client.get(f"/api/v1/admin/reports/{key}", headers=staff_headers(admin))
        assert r.status_code == 200, (key, r.text)
    x = client.get("/api/v1/admin/reports/referrers_performance/export?format=xlsx", headers=staff_headers(admin))
    assert x.status_code == 200 and x.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(x.content))
    values = [c.value for row in wb.active.iter_rows() for c in row if isinstance(c.value, str)]
    assert not any(v.startswith("=") for v in values) and wb.active.sheet_view.rightToLeft
    c = client.get("/api/v1/admin/reports/referrers_performance/export?format=csv", headers=staff_headers(admin))
    assert c.content.startswith("﻿".encode()) and b"'=HYPERLINK" in c.content
    assert client.get("/api/v1/admin/reports/points_value/export", headers=staff_headers(op)).status_code == 403  # no reports.export
    s1 = client.post("/api/v1/admin/reports/points_value/snapshots", json={"note": "گزارش ماهانه"}, headers=staff_headers(admin)).json()
    s2 = client.post("/api/v1/admin/reports/points_value/snapshots", json={}, headers=staff_headers(admin)).json()
    assert (s1["version"], s2["version"]) == (1, 2) and client.get(s1["xlsx_url"]).status_code == 200
    lst = client.get("/api/v1/admin/reports/snapshots?report_key=points_value", headers=staff_headers(admin)).json()
    assert lst["total"] == 2


def test_broadcast_segment(client, db):
    op = make_staff(db, "operator")
    a = make_user(db, sn="1000000001", mobile="09120000001")
    a.tier = "high_cost"
    b = make_user(db, sn="1000000002", mobile="09120000002")
    b.tier = "low_cost"
    make_user(db, sn="1000000003", mobile="09120000003")
    db.commit()
    seg = {"tiers": ["high_cost", "none"]}
    assert client.post("/api/v1/admin/notifications/broadcasts/preview", json=seg, headers=staff_headers(op)).json()["recipients"] == 2
    r = client.post("/api/v1/admin/notifications/broadcasts", json={"title": "هفتهٔ صرفه‌جویی", "body": "راهکارهای جدید",
                                                                    "channels": ["in_app", "sms"], "segment": seg, "route": "/learn"},
                    headers=staff_headers(op))
    assert r.status_code == 201 and r.json()["recipients"] == 2
    db.rollback()
    assert db.scalar(select(func.count()).select_from(Notification).where(Notification.kind == "broadcast")) == 2
    assert db.scalar(select(func.count()).select_from(SmsOutbox).where(SmsOutbox.purpose == "broadcast")) == 2
    assert client.get("/api/v1/app/notifications", headers=user_headers(b)).json()["total"] == 0


def test_survey_points_once_and_validation(client, db):
    op = make_staff(db, "super_admin")
    u = make_user(db)
    db.commit()
    body = {"title": "رضایت از طرح گازیار", "questions": [
        {"kind": "rating", "text": "چقدر از طرح راضی هستید؟"},
        {"kind": "single", "text": "از کجا با طرح آشنا شدید؟", "options": ["پیامک", "دوستان", "سفیر"]},
        {"kind": "text", "text": "پیشنهاد شما", "required": False}]}
    s = client.post("/api/v1/admin/surveys", json=body, headers=staff_headers(op)).json()
    assert client.get("/api/v1/app/surveys", headers=user_headers(u)).json() == []
    client.post(f"/api/v1/admin/surveys/{s['id']}/status", json={"status": "active"}, headers=staff_headers(op))
    sv = client.get("/api/v1/app/surveys", headers=user_headers(u)).json()[0]
    assert sv["points"] == 10
    q = {x["kind"]: x["id"] for x in sv["questions"]}
    bad = client.post(f"/api/v1/app/surveys/{s['id']}/answers", json={"answers": {str(q["rating"]): 9}}, headers=user_headers(u))
    assert bad.status_code == 422 and set(bad.json()["error"]["details"]["fields"]) == {str(q["rating"]), str(q["single"])}
    ok = client.post(f"/api/v1/app/surveys/{s['id']}/answers", json={"answers": {str(q["rating"]): 5, str(q["single"]): "سفیر"}},
                     headers=user_headers(u))
    assert ok.status_code == 201 and ok.json()["answered"]
    assert client.post(f"/api/v1/app/surveys/{s['id']}/answers", json={"answers": {str(q["rating"]): 4, str(q["single"]): "سفیر"}},
                       headers=user_headers(u)).status_code == 409
    db.rollback()
    assert db.scalar(select(func.sum(PointsLedger.amount)).where(PointsLedger.event_type == "survey_completed")) == 10
    res = client.get(f"/api/v1/admin/surveys/{s['id']}/results", headers=staff_headers(op)).json()
    assert res["responses"] == 1 and res["questions"][0]["average"] == 5


def test_stage3_permissions(client, db):
    amb = make_staff(db, "ambassador")
    op = make_staff(db, "operator")
    db.commit()
    for path in ("/api/v1/admin/finance/redemptions", "/api/v1/admin/finance/summary", "/api/v1/admin/store/items"):
        assert client.get(path, headers=staff_headers(amb)).status_code == 403
    assert client.get("/api/v1/admin/finance/redemptions", headers=staff_headers(op)).status_code == 403
    assert client.post("/api/v1/admin/notifications/broadcasts/preview", json={}, headers=staff_headers(amb)).status_code == 403


