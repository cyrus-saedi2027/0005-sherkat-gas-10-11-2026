"""Excel import of the authorised subscriber list + 3-year consumption history (R01, R04).

Two-step and idempotent: upload → validate (dry run, nothing written except the
batch report) → commit (upsert). Re-importing the same file changes nothing.

Accepted layouts (headers are matched after Persian normalisation, Persian or English):
  sheet «مشترکین»:  شماره اشتراک | شماره موبایل | نام | کد ناحیه | آدرس | شماره بدنه کنتور
  sheet «مصرف» long: شماره اشتراک | سال | ماه | مصرف
  sheet «مصرف» wide: شماره اشتراک | 1402-01 | 1402-02 | …
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import insert, select, update
from sqlalchemy.orm import Session

from app.core import clock, crypto
from app.core.errors import AppError
from app.core.normalize import normalize_mobile, normalize_subscription, normalize_text, to_ascii_digits
from app.modules.regions.models import Region
from app.modules.subscribers import service as svc
from app.modules.subscribers.models import ConsumptionHistory, ImportBatch, Subscriber

MAX_ROWS = 200_000
MAX_ERRORS = 500

SUB_COLS: dict[str, tuple[str, ...]] = {
    "subscription_no": ("شماره اشتراک", "اشتراک", "subscription_no", "subscription"),
    "mobile": ("شماره موبایل", "موبایل", "تلفن همراه", "mobile"),
    "full_name": ("نام", "نام و نام خانوادگی", "نام مشترک", "full_name", "name"),
    "region_code": ("کد ناحیه", "ناحیه", "کد منطقه", "region_code", "region"),
    "address": ("آدرس", "نشانی", "address"),
    "meter_serial": ("شماره بدنه کنتور", "سریال کنتور", "شماره کنتور", "meter_serial"),
}
CONS_COLS: dict[str, tuple[str, ...]] = {
    "subscription_no": SUB_COLS["subscription_no"],
    "jyear": ("سال", "year", "jyear"),
    "jmonth": ("ماه", "month", "jmonth"),
    "volume_m3": ("مصرف", "مصرف (مترمکعب)", "مصرف مترمکعب", "volume", "volume_m3"),
}
SUB_SHEETS = ("مشترکین", "مشترکان", "subscribers")
CONS_SHEETS = ("مصرف", "سوابق مصرف", "consumption")
PERIOD_HEADER = re.compile(r"^(\d{4})[-/](\d{1,2})$")


@dataclass
class Parsed:
    subscribers: list[dict[str, Any]] = field(default_factory=list)
    history: list[dict[str, Any]] = field(default_factory=list)
    errors: list[dict[str, Any]] = field(default_factory=list)
    rows: int = 0

    def err(self, sheet: str, row: int, column: str, message: str) -> None:
        if len(self.errors) < MAX_ERRORS:
            self.errors.append({"sheet": sheet, "row": row, "column": column, "message": message})


def _norm(h: Any) -> str:
    return normalize_text(str(h or "")).replace("_", " ").strip()


def _map_headers(header: list[Any], spec: dict[str, tuple[str, ...]]) -> dict[str, int]:
    norm = [_norm(h) for h in header]
    out: dict[str, int] = {}
    for key, aliases in spec.items():
        wanted = {_norm(a) for a in aliases}
        for i, h in enumerate(norm):
            if h in wanted:
                out[key] = i
                break
    return out


def _find_sheet(wb: Any, names: tuple[str, ...]) -> Any:
    wanted = {_norm(n) for n in names}
    for ws in wb.worksheets:
        if _norm(ws.title) in wanted:
            return ws
    return None


def _num(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, int | float):
        return float(v)
    try:
        return float(to_ascii_digits(str(v)).replace(",", "").replace("٬", "").replace("٫", ".").strip())
    except ValueError:
        return None


def parse(raw: bytes, db: Session) -> Parsed:
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise AppError("bad_file", "فایل اکسل قابل خواندن نیست. فقط فایل ‎.xlsx‎ پذیرفته می‌شود.", 422) from e
    p = Parsed()
    regions = {normalize_text(r.code): r.id for r in db.scalars(select(Region)).all()}
    ws_sub = _find_sheet(wb, SUB_SHEETS)
    ws_cons = _find_sheet(wb, CONS_SHEETS)
    if ws_sub is None and ws_cons is None:
        if len(wb.worksheets) == 1:
            ws_sub = wb.worksheets[0]
        else:
            raise AppError("bad_file", "برگه‌ای با نام «مشترکین» یا «مصرف» پیدا نشد.", 422)

    seen_subs: set[str] = set()
    if ws_sub is not None:
        it = ws_sub.iter_rows(values_only=True)
        header = list(next(it, []) or [])
        cols = _map_headers(header, SUB_COLS)
        missing = [k for k in ("subscription_no", "mobile") if k not in cols]
        if missing:
            raise AppError("bad_file", "ستون‌های «شماره اشتراک» و «شماره موبایل» در برگهٔ مشترکین الزامی‌اند.", 422)
        wide_periods = {i: PERIOD_HEADER.match(to_ascii_digits(str(h or "")).strip()) for i, h in enumerate(header)}
        wide_periods = {i: m for i, m in wide_periods.items() if m}
        for rn, row in enumerate(it, start=2):
            if not row or all(c in (None, "") for c in row):
                continue
            p.rows += 1
            if p.rows > MAX_ROWS:
                raise AppError("too_many_rows", f"حداکثر {MAX_ROWS} ردیف در هر فایل مجاز است.", 422)

            def cell(key: str, row: tuple = row) -> Any:
                i = cols.get(key)
                return row[i] if i is not None and i < len(row) else None

            sn = normalize_subscription(str(cell("subscription_no") or ""))
            if not sn:
                p.err(ws_sub.title, rn, "شماره اشتراک", "شمارهٔ اشتراک معتبر نیست.")
                continue
            if sn in seen_subs:
                p.err(ws_sub.title, rn, "شماره اشتراک", "شمارهٔ اشتراک در فایل تکراری است.")
                continue
            m = normalize_mobile(str(cell("mobile") or ""))
            if not m:
                p.err(ws_sub.title, rn, "شماره موبایل", "شمارهٔ موبایل معتبر نیست.")
                continue
            rid = None
            rc = cell("region_code")
            if rc not in (None, ""):
                rid = regions.get(normalize_text(str(rc)))
                if rid is None:
                    p.err(ws_sub.title, rn, "کد ناحیه", f"کد ناحیهٔ «{rc}» تعریف نشده است.")
                    continue
            seen_subs.add(sn)
            p.subscribers.append({"subscription_no": sn, "mobile": m,
                                  "full_name": (normalize_text(str(cell("full_name") or "")) or None),
                                  "region_id": rid, "address": (str(cell("address") or "").strip()[:500] or None),
                                  "meter_serial": (str(cell("meter_serial") or "").strip()[:40] or None)})
            for i, mt in wide_periods.items():  # a single-sheet file may carry history columns too
                v = _num(row[i] if i < len(row) else None)
                if v is not None:
                    _add_hist(p, ws_sub.title, rn, sn, int(mt.group(1)), int(mt.group(2)), v)

    if ws_cons is not None:
        it = ws_cons.iter_rows(values_only=True)
        header = list(next(it, []) or [])
        cols = _map_headers(header, CONS_COLS)
        wide = {i: PERIOD_HEADER.match(to_ascii_digits(str(h or "")).strip()) for i, h in enumerate(header)}
        wide = {i: m for i, m in wide.items() if m}
        if "subscription_no" not in cols:
            raise AppError("bad_file", "ستون «شماره اشتراک» در برگهٔ مصرف الزامی است.", 422)
        is_long = all(k in cols for k in ("jyear", "jmonth", "volume_m3"))
        if not is_long and not wide:
            raise AppError("bad_file", "برگهٔ مصرف باید ستون‌های «سال، ماه، مصرف» یا ستون‌هایی مثل «1402-01» داشته باشد.", 422)
        for rn, row in enumerate(it, start=2):
            if not row or all(c in (None, "") for c in row):
                continue
            p.rows += 1
            if p.rows > MAX_ROWS:
                raise AppError("too_many_rows", f"حداکثر {MAX_ROWS} ردیف در هر فایل مجاز است.", 422)
            sn = normalize_subscription(str(row[cols["subscription_no"]] or ""))
            if not sn:
                p.err(ws_cons.title, rn, "شماره اشتراک", "شمارهٔ اشتراک معتبر نیست.")
                continue
            if is_long:
                y, mth, v = (_num(row[cols[k]]) for k in ("jyear", "jmonth", "volume_m3"))
                if y is None or mth is None or v is None:
                    p.err(ws_cons.title, rn, "سال/ماه/مصرف", "سال، ماه و مصرف باید عدد باشند.")
                    continue
                _add_hist(p, ws_cons.title, rn, sn, int(y), int(mth), v)
            else:
                for i, mt in wide.items():
                    v = _num(row[i] if i < len(row) else None)
                    if v is not None:
                        _add_hist(p, ws_cons.title, rn, sn, int(mt.group(1)), int(mt.group(2)), v)
    wb.close()

    # History must reference a subscriber in this file or already in the database.
    if p.history:
        known = seen_subs | set(db.scalars(select(Subscriber.subscription_no).where(
            Subscriber.subscription_no.in_({h["subscription_no"] for h in p.history}))).all())
        kept = []
        for h in p.history:
            if h["subscription_no"] in known:
                kept.append(h)
            else:
                p.err(h["sheet"], h["row"], "شماره اشتراک", "این اشتراک نه در فایل هست و نه در سامانه.")
        p.history = kept
    return p


def _add_hist(p: Parsed, sheet: str, rn: int, sn: str, y: int, m: int, v: float) -> None:
    if not (1300 <= y <= 1500 and 1 <= m <= 12):
        p.err(sheet, rn, "سال/ماه", "سال یا ماه جلالی معتبر نیست.")
        return
    if v < 0 or v > 1_000_000:
        p.err(sheet, rn, "مصرف", "مقدار مصرف خارج از محدودهٔ مجاز است.")
        return
    p.history.append({"subscription_no": sn, "jyear": y, "jmonth": m, "volume_m3": float(v), "sheet": sheet, "row": rn})


def summarize(db: Session, p: Parsed) -> dict[str, int]:
    existing = set(db.scalars(select(Subscriber.subscription_no).where(
        Subscriber.subscription_no.in_([s["subscription_no"] for s in p.subscribers]))).all()) if p.subscribers else set()
    return {"subscribers_new": sum(1 for s in p.subscribers if s["subscription_no"] not in existing),
            "subscribers_updated": sum(1 for s in p.subscribers if s["subscription_no"] in existing),
            "history_rows": len(p.history), "errors": len(p.errors), "rows": p.rows}


def commit(db: Session, p: Parsed, batch: ImportBatch) -> dict[str, int]:
    by_sn = {s.subscription_no: s for s in db.scalars(select(Subscriber).where(
        Subscriber.subscription_no.in_([x["subscription_no"] for x in p.subscribers] or [""]))).all()}
    created = updated = 0
    for row in p.subscribers:
        enc, bidx = crypto.encrypt(row["mobile"]), svc.mobile_bidx(row["mobile"])
        s = by_sn.get(row["subscription_no"])
        if s is None:
            s = Subscriber(subscription_no=row["subscription_no"], mobile_enc=enc, mobile_bidx=bidx,
                           full_name=row["full_name"], region_id=row["region_id"], address=row["address"],
                           meter_serial=row["meter_serial"], import_batch_id=batch.id)
            db.add(s)
            by_sn[s.subscription_no] = s
            created += 1
        else:
            if s.mobile_bidx != bidx:
                s.mobile_enc, s.mobile_bidx = enc, bidx
            for k in ("full_name", "region_id", "address", "meter_serial"):
                if row[k] is not None:
                    setattr(s, k, row[k])
            s.import_batch_id, s.is_active = batch.id, True
            updated += 1
    db.flush()
    if p.history:
        missing = {h["subscription_no"] for h in p.history} - set(by_sn)
        if missing:
            for s in db.scalars(select(Subscriber).where(Subscriber.subscription_no.in_(missing))).all():
                by_sn[s.subscription_no] = s
        sids = {by_sn[h["subscription_no"]].id for h in p.history if h["subscription_no"] in by_sn}
        existing = {(r.subscriber_id, r.jyear, r.jmonth): r.id for r in db.execute(
            select(ConsumptionHistory.id, ConsumptionHistory.subscriber_id, ConsumptionHistory.jyear, ConsumptionHistory.jmonth)
            .where(ConsumptionHistory.subscriber_id.in_(sids))).all()}
        inserts, updates, dedup = [], [], {}
        now = clock.now()
        for h in p.history:
            s = by_sn.get(h["subscription_no"])
            if s is None:
                continue
            dedup[(s.id, h["jyear"], h["jmonth"])] = h["volume_m3"]  # last value in the file wins
        for (sid, y, m), v in dedup.items():
            rid = existing.get((sid, y, m))
            if rid is None:
                inserts.append({"subscriber_id": sid, "jyear": y, "jmonth": m, "volume_m3": v, "import_batch_id": batch.id,
                                "created_at": now, "updated_at": now})
            else:
                updates.append({"id": rid, "volume_m3": v, "import_batch_id": batch.id, "updated_at": now})
        for i in range(0, len(inserts), 5000):
            db.execute(insert(ConsumptionHistory), inserts[i:i + 5000])
        if updates:
            db.execute(update(ConsumptionHistory), updates)
        svc.refresh_baselines(db, list(sids))
    return {"subscribers_created": created, "subscribers_updated": updated, "history_rows": len(p.history)}


def template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "مشترکین"
    ws.sheet_view.rightToLeft = True
    ws.append(["شماره اشتراک", "شماره موبایل", "نام و نام خانوادگی", "کد ناحیه", "آدرس", "شماره بدنه کنتور"])
    ws.append(["1234567890", "09121234567", "نمونه مشترک", "R2-Z1", "رشت، ...", "MTR-0001"])
    ws2 = wb.create_sheet("مصرف")
    ws2.sheet_view.rightToLeft = True
    ws2.append(["شماره اشتراک", "سال", "ماه", "مصرف (مترمکعب)"])
    for y in (1402, 1403, 1404):
        for m in range(1, 13):
            ws2.append(["1234567890", y, m, 300 if m in (9, 10, 11, 12, 1, 2) else 80])
    for sheet in (ws, ws2):
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
