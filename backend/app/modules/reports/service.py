"""Organisational reporting engine (R31, R41).

Every report is a function returning a Table (columns + rows + summary). The same
Table renders as JSON (panel + charts), Excel (RTL sheet) and CSV (UTF-8 BOM so
Excel opens Persian correctly). Snapshots freeze the params + files as a version.
"""

from __future__ import annotations

import csv
import io
from collections import defaultdict
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core import clock, jalali
from app.core import settings_registry as reg
from app.core.errors import AppError
from app.modules.regions.models import Region
from app.modules.regions.service import descendants

MAX_ROWS = 20_000


@dataclass
class Params:
    region_id: int | None = None
    date_from: date | None = None
    date_to: date | None = None
    tier: str | None = None
    expert_id: int | None = None

    def bounds(self) -> tuple[datetime, datetime]:
        today = jalali.to_tehran(clock.now()).date()
        df = self.date_from or today - timedelta(days=364)
        dt = self.date_to or today
        return (datetime.combine(df, time(0), tzinfo=jalali.TEHRAN),
                datetime.combine(dt + timedelta(days=1), time(0), tzinfo=jalali.TEHRAN))

    def regions(self, db: Session) -> set[int] | None:
        return descendants(db, self.region_id) if self.region_id else None


@dataclass
class Column:
    key: str
    title: str
    type: str = "text"  # text | int | float | pct | rial | m3 | date


@dataclass
class Table:
    title: str
    columns: list[Column]
    rows: list[dict[str, Any]]
    summary: dict[str, Any] = field(default_factory=dict)
    chart: dict[str, Any] | None = None


@dataclass
class ReportDef:
    key: str
    title: str
    description: str
    fn: Callable[[Session, Params], Table]


REPORTS: dict[str, ReportDef] = {}


def report(key: str, title: str, description: str):
    def deco(fn: Callable[[Session, Params], Table]):
        REPORTS[key] = ReportDef(key, title, description, fn)
        return fn

    return deco


def run(db: Session, key: str, params: Params) -> Table:
    d = REPORTS.get(key)
    if d is None:
        raise AppError("not_found", "گزارش پیدا نشد.", 404)
    t = d.fn(db, params)
    if len(t.rows) > MAX_ROWS:
        t.rows = t.rows[:MAX_ROWS]
        t.summary["truncated_to"] = MAX_ROWS
    return t


def _region_names(db: Session) -> dict[int, str]:
    return dict(db.execute(select(Region.id, Region.name)).all())


# ───────────────────────── reports ─────────────────────────

@report("consumption_comparison", "مقایسهٔ مصرف قبل و بعد", "مصرف واقعی هر مشترک در برابر میانگین دورهٔ مشابه سه سال قبل و درصد کاهش")
def consumption_comparison(db: Session, p: Params) -> Table:
    from app.modules.readings.models import MeterReading
    from app.modules.subscribers.models import Subscriber
    from app.modules.users.models import User

    start, end = p.bounds()
    stmt = (select(Subscriber.subscription_no, Subscriber.full_name, Subscriber.region_id, User.tier,
                   func.count(MeterReading.id), func.sum(MeterReading.baseline_m3), func.sum(MeterReading.consumption_m3))
            .join(MeterReading, MeterReading.subscriber_id == Subscriber.id)
            .join(User, User.id == MeterReading.user_id)
            .where(MeterReading.status == "approved", MeterReading.is_current.is_(True),
                   MeterReading.consumption_m3.is_not(None), MeterReading.baseline_m3.is_not(None),
                   MeterReading.reviewed_at >= start, MeterReading.reviewed_at < end)
            .group_by(Subscriber.subscription_no, Subscriber.full_name, Subscriber.region_id, User.tier))
    if (rs := p.regions(db)) is not None:
        stmt = stmt.where(Subscriber.region_id.in_(rs))
    if p.tier:
        stmt = stmt.where(User.tier == p.tier)
    names = _region_names(db)
    rows, tb, ta = [], 0.0, 0.0
    by_region: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0])
    for sn, name, rid, tier, n, base, act in db.execute(stmt).all():
        base, act = float(base or 0), float(act or 0)
        red = round((base - act) / base * 100, 1) if base else None
        rows.append({"subscription_no": sn, "name": name, "region": names.get(rid, "—"), "tier": tier or "—", "periods": int(n),
                     "baseline_m3": round(base, 1), "actual_m3": round(act, 1), "saved_m3": round(base - act, 1),
                     "reduction_pct": red})
        tb, ta = tb + base, ta + act
        agg = by_region[names.get(rid, "—")]
        agg[0] += base
        agg[1] += act
    rows.sort(key=lambda r: -(r["reduction_pct"] or -1e9))
    return Table("مقایسهٔ مصرف قبل و بعد",
                 [Column("subscription_no", "شماره اشتراک"), Column("name", "نام"), Column("region", "ناحیه"),
                  Column("tier", "گروه"), Column("periods", "تعداد دوره", "int"), Column("baseline_m3", "میانگین سه سال قبل", "m3"),
                  Column("actual_m3", "مصرف واقعی", "m3"), Column("saved_m3", "صرفه‌جویی", "m3"),
                  Column("reduction_pct", "درصد کاهش", "pct")],
                 rows, {"subscribers": len(rows), "baseline_m3": round(tb, 1), "actual_m3": round(ta, 1),
                        "saved_m3": round(tb - ta, 1), "reduction_pct": round((tb - ta) / tb * 100, 1) if tb else None},
                 {"type": "bar", "categories": list(by_region), "series": [
                     {"name": "میانگین سه سال قبل", "data": [round(v[0], 1) for v in by_region.values()]},
                     {"name": "مصرف واقعی", "data": [round(v[1], 1) for v in by_region.values()]}]})


@report("consumers_low_high", "مشترکان کم‌مصرف و پرمصرف", "رتبه‌بندی مشترکان بر اساس مصرف مبنای سالانه")
def consumers_low_high(db: Session, p: Params) -> Table:
    from app.modules.subscribers.models import Subscriber
    from app.modules.users.models import User

    threshold = reg.get_int(db, "eligibility.threshold_m3")
    stmt = select(Subscriber, User.id).outerjoin(User, User.subscriber_id == Subscriber.id).where(
        Subscriber.baseline_annual_m3.is_not(None))
    if (rs := p.regions(db)) is not None:
        stmt = stmt.where(Subscriber.region_id.in_(rs))
    names = _region_names(db)
    rows = []
    for s, uid in db.execute(stmt.order_by(Subscriber.baseline_annual_m3.desc())).all():
        band = "پرمصرف" if s.baseline_annual_m3 >= threshold * 1.5 else ("متوسط" if s.baseline_annual_m3 >= threshold else "کم‌مصرف")
        rows.append({"subscription_no": s.subscription_no, "name": s.full_name, "region": names.get(s.region_id, "—"),
                     "baseline_annual_m3": round(s.baseline_annual_m3, 1), "band": band, "registered": "بله" if uid else "خیر"})
    counts = defaultdict(int)
    for r in rows:
        counts[r["band"]] += 1
    return Table("مشترکان کم‌مصرف و پرمصرف",
                 [Column("subscription_no", "شماره اشتراک"), Column("name", "نام"), Column("region", "ناحیه"),
                  Column("baseline_annual_m3", "مصرف مبنای سالانه", "m3"), Column("band", "گروه مصرف"), Column("registered", "ثبت‌نام")],
                 rows, {"threshold_m3": threshold, **{f"count_{k}": v for k, v in counts.items()}},
                 {"type": "donut", "categories": list(counts), "series": [{"name": "تعداد", "data": list(counts.values())}]})


@report("referrers_performance", "عملکرد معرف‌ها و نرخ تبدیل", "تعداد معرفی، معرفی معتبر، مشکوک و نرخ تبدیل هر معرف")
def referrers_performance(db: Session, p: Params) -> Table:
    from app.modules.referrals.models import Referral
    from app.modules.users.models import User

    start, end = p.bounds()
    stmt = (select(Referral.referrer_user_id, Referral.status, func.count()).where(Referral.created_at >= start,
                                                                                   Referral.created_at < end)
            .group_by(Referral.referrer_user_id, Referral.status))
    agg: dict[int, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for uid, st, n in db.execute(stmt).all():
        agg[uid][st] += int(n)
    users = {u.id: u for u in db.scalars(select(User).where(User.id.in_(list(agg) or [0]))).all()}
    rows = []
    for uid, c in agg.items():
        total = sum(c.values())
        u = users.get(uid)
        rows.append({"user_id": uid, "name": " ".join(x for x in ((u.first_name, u.last_name) if u else ()) if x) or "—",
                     "total": total, "rewarded": c.get("rewarded", 0), "pending": c.get("pending", 0),
                     "suspicious": c.get("suspicious", 0), "capped": c.get("capped", 0), "rejected": c.get("rejected", 0),
                     "conversion_pct": round(c.get("rewarded", 0) / total * 100, 1) if total else 0})
    rows.sort(key=lambda r: (-r["rewarded"], -r["total"]))
    tot = sum(r["total"] for r in rows)
    ok = sum(r["rewarded"] for r in rows)
    return Table("عملکرد معرف‌ها",
                 [Column("user_id", "شناسهٔ کاربر", "int"), Column("name", "نام معرف"), Column("total", "کل معرفی", "int"),
                  Column("rewarded", "معرفی معتبر", "int"), Column("pending", "در انتظار", "int"),
                  Column("suspicious", "مشکوک", "int"), Column("capped", "خارج از سقف", "int"), Column("rejected", "ردشده", "int"),
                  Column("conversion_pct", "نرخ تبدیل", "pct")],
                 rows, {"referrals": tot, "valid_referrals": ok, "conversion_pct": round(ok / tot * 100, 1) if tot else 0})


@report("cost_per_action", "هزینهٔ هر اقدام", "هزینهٔ پرونده‌های تکمیل‌شده به تفکیک گروه، سهم طرح و مشترک و هزینه به ازای هر مترمکعب صرفه‌جویی")
def cost_per_action(db: Session, p: Params) -> Table:
    from app.modules.cases.models import Case, CaseEstimate
    from app.modules.readings.models import MeterReading

    start, end = p.bounds()
    stmt = select(Case).where(Case.status == "completed", Case.completed_at >= start, Case.completed_at < end)
    if (rs := p.regions(db)) is not None:
        stmt = stmt.where(Case.region_id.in_(rs))
    if p.tier:
        stmt = stmt.where(Case.tier == p.tier)
    if p.expert_id:
        stmt = stmt.where(Case.ambassador_id == p.expert_id)
    agg: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for c in db.scalars(stmt).all():
        est = db.scalar(select(CaseEstimate).where(CaseEstimate.case_id == c.id, CaseEstimate.status == "accepted")
                        .order_by(CaseEstimate.id.desc()))
        saved = db.scalar(select(func.coalesce(func.sum(MeterReading.baseline_m3 - MeterReading.consumption_m3), 0)).where(
            MeterReading.user_id == c.user_id, MeterReading.status == "approved", MeterReading.is_current.is_(True),
            MeterReading.consumption_m3 < MeterReading.baseline_m3, MeterReading.reviewed_at >= (c.completed_at or start))) or 0
        a = agg[c.tier or "—"]
        a["cases"] += 1
        a["total_cost"] += est.total_cost_rial if est else 0
        a["program_share"] += est.program_share_rial if est else 0
        a["subscriber_share"] += est.subscriber_share_rial if est else 0
        a["saved_m3"] += float(saved)
    rows = []
    for tier, a in agg.items():
        rows.append({"tier": tier, "cases": int(a["cases"]), "total_cost_rial": int(a["total_cost"]),
                     "program_share_rial": int(a["program_share"]), "subscriber_share_rial": int(a["subscriber_share"]),
                     "avg_cost_rial": int(a["total_cost"] / a["cases"]) if a["cases"] else 0,
                     "saved_m3": round(a["saved_m3"], 1),
                     "cost_per_m3_rial": int(a["program_share"] / a["saved_m3"]) if a["saved_m3"] > 0 else None})
    return Table("هزینهٔ هر اقدام",
                 [Column("tier", "گروه"), Column("cases", "پرونده‌های تکمیل‌شده", "int"), Column("total_cost_rial", "کل هزینه", "rial"),
                  Column("program_share_rial", "سهم طرح", "rial"), Column("subscriber_share_rial", "سهم مشترک", "rial"),
                  Column("avg_cost_rial", "میانگین هزینهٔ هر پرونده", "rial"), Column("saved_m3", "گاز صرفه‌جویی‌شده", "m3"),
                  Column("cost_per_m3_rial", "هزینهٔ طرح به ازای هر مترمکعب", "rial")],
                 rows, {"cases": sum(r["cases"] for r in rows), "program_share_rial": sum(r["program_share_rial"] for r in rows)})


@report("points_value", "ارزش امتیازها و نرخ تبدیل پاداش", "امتیاز صادرشده به تفکیک رویداد، ارزش ریالی و میزان استفاده‌شده")
def points_value(db: Session, p: Params) -> Table:
    from app.modules.points.models import PointsLedger
    from app.modules.points.service import event_title

    start, end = p.bounds()
    rate = reg.get_int(db, "finance.point_rial_rate")
    stmt = (select(PointsLedger.event_type, func.count(), func.sum(PointsLedger.amount))
            .where(PointsLedger.created_at >= start, PointsLedger.created_at < end, PointsLedger.status != "reversed")
            .group_by(PointsLedger.event_type))
    rows, issued, spent = [], 0, 0
    for et, n, total in db.execute(stmt).all():
        total = int(total or 0)
        if total > 0 and et not in ("redemption_refund",):
            issued += total
        if et == "redemption":
            spent += -total
        rows.append({"event_type": et, "title": event_title(db, et), "entries": int(n), "points": total, "rial_value": total * rate})
    rows.sort(key=lambda r: -abs(r["points"]))
    return Table("ارزش امتیازها",
                 [Column("title", "رویداد"), Column("entries", "تعداد", "int"), Column("points", "امتیاز", "int"),
                  Column("rial_value", "ارزش ریالی", "rial")],
                 rows, {"issued_points": issued, "issued_rial": issued * rate, "redeemed_points": spent,
                        "conversion_pct": round(spent / issued * 100, 1) if issued else 0, "point_rial_rate": rate},
                 {"type": "bar", "categories": [r["title"] for r in rows if r["points"] > 0],
                  "series": [{"name": "امتیاز", "data": [r["points"] for r in rows if r["points"] > 0]}]})


@report("readings_review", "عملکرد بررسی قرائت‌ها", "تعداد تأیید، رد و نیاز به اصلاح به تفکیک کارشناس")
def readings_review(db: Session, p: Params) -> Table:
    from app.modules.readings.models import MeterReading
    from app.modules.staff.models import StaffUser

    start, end = p.bounds()
    stmt = (select(MeterReading.reviewer_id, MeterReading.status, func.count())
            .where(MeterReading.reviewed_at >= start, MeterReading.reviewed_at < end, MeterReading.reviewer_id.is_not(None))
            .group_by(MeterReading.reviewer_id, MeterReading.status))
    if p.expert_id:
        stmt = stmt.where(MeterReading.reviewer_id == p.expert_id)
    agg: dict[int, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for rid, st, n in db.execute(stmt).all():
        agg[rid][st] += int(n)
    names = dict(db.execute(select(StaffUser.id, StaffUser.full_name)).all())
    rows = [{"reviewer": names.get(rid, "—"), "approved": c.get("approved", 0), "needs_correction": c.get("needs_correction", 0),
             "rejected": c.get("rejected", 0), "total": sum(c.values())} for rid, c in agg.items()]
    return Table("عملکرد بررسی قرائت‌ها",
                 [Column("reviewer", "کارشناس"), Column("approved", "تأیید", "int"), Column("needs_correction", "نیاز به اصلاح", "int"),
                  Column("rejected", "رد", "int"), Column("total", "جمع", "int")], rows,
                 {"total": sum(r["total"] for r in rows)})


@report("tickets_satisfaction", "تیکت‌ها و رضایت", "وضعیت تیکت‌ها و میانگین امتیاز رضایت به تفکیک واحد")
def tickets_satisfaction(db: Session, p: Params) -> Table:
    from app.modules.tickets.models import Ticket

    start, end = p.bounds()
    stmt = (select(Ticket.assigned_unit, func.count(), func.avg(Ticket.rating), func.count(Ticket.rating))
            .where(Ticket.created_at >= start, Ticket.created_at < end).group_by(Ticket.assigned_unit))
    unit_fa = {"support": "پشتیبانی", "technical": "فنی", "finance": "مالی", "ambassador": "سفیران"}
    rows = [{"unit": unit_fa.get(u, u), "tickets": int(n), "avg_rating": round(float(avg), 2) if avg is not None else None,
             "rated": int(rated)} for u, n, avg, rated in db.execute(stmt).all()]
    return Table("تیکت‌ها و رضایت", [Column("unit", "واحد"), Column("tickets", "تعداد تیکت", "int"),
                                      Column("avg_rating", "میانگین رضایت (از ۵)", "float"), Column("rated", "تعداد امتیازدهی", "int")],
                 rows, {"tickets": sum(r["tickets"] for r in rows)})


# ───────────────────────── exports ─────────────────────────

def to_xlsx(t: Table, params: Params) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = t.title[:31]
    ws.sheet_view.rightToLeft = True
    ws.append([t.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"تاریخ تهیه: {jalali.format_date(clock.now())}",
               f"بازه: {jalali.format_date(params.bounds()[0])} تا {jalali.format_date(params.bounds()[1] - timedelta(days=1))}"])
    ws.append([])
    ws.append([c.title for c in t.columns])
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1546C8")
        cell.alignment = Alignment(horizontal="center")
    for r in t.rows:
        ws.append([_safe_cell(r.get(c.key)) for c in t.columns])
    fmt = {"rial": "#,##0", "int": "#,##0", "m3": "#,##0.0", "float": "0.00", "pct": "0.0"}
    for i, c in enumerate(t.columns, start=1):
        letter = ws.cell(row=4, column=i).column_letter
        ws.column_dimensions[letter].width = max(14, len(c.title) + 6)
        if c.type in fmt:
            for row in ws.iter_rows(min_row=5, min_col=i, max_col=i):
                row[0].number_format = fmt[c.type]
    if t.summary:
        s = wb.create_sheet("خلاصه")
        s.sheet_view.rightToLeft = True
        for k, v in t.summary.items():
            s.append([k, v])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _safe_cell(v: Any) -> Any:
    # openpyxl turns any string starting with "=" into a live formula; user-supplied text
    # (names, titles) must never execute in the finance team's Excel (formula injection).
    return ("'" + v) if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


def to_csv(t: Table) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c.title for c in t.columns])
    for r in t.rows:
        # Guard against CSV/formula injection when the file is opened in Excel.
        w.writerow([_safe_cell(r.get(c.key)) for c in t.columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")
